"""Documents read locally, web search with a fake network, token economy. Fake data only, no network."""

import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("MNEMA_HOME", tempfile.mkdtemp())

from mnema import docs, retrieve, web  # noqa: E402
from mnema.agent import profile as P  # noqa: E402
from mnema.chat import attach, economy as EC, engine as CH, providers as PR  # noqa: E402
from mnema.memory import init_vault  # noqa: E402
from mnema.vault import Vault  # noqa: E402
from test_core import fill  # noqa: E402
from test_v3 import FakeStream  # noqa: E402


def make_pdf(path: Path, pages: list[str]) -> None:
    """A tiny valid PDF with one line of Latin text per page (Helvetica)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for t in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({t}) Tj ET".encode()
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {len(objs)} 0 R >>"
        )
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode()
    path.write_bytes(out)


class Base(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class TestDocs(Base):
    def test_pdf_pages(self):
        p = self.dir / "offer.pdf"
        make_pdf(p, ["Pilot price is 120 thousand", "Delivery in two weeks"])
        text, info = docs.extract(p)
        self.assertEqual(info["pages"], 2)
        self.assertIn("[стр. 2]", text)
        self.assertIn("Pilot price", text)

    def test_xlsx_all_sheets(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "Бюджет"
        wb.active.append(["Статья", "Сумма"])
        wb.active.append(["Серверы", 4500])
        wb.create_sheet("План").append(["Неделя 1", "созвоны"])
        p = self.dir / "b.xlsx"
        wb.save(p)
        text, info = docs.extract(p)
        self.assertEqual(info["sheets"], 2)
        self.assertIn("## Лист «Бюджет»", text)
        self.assertIn("Серверы | 4500", text)
        self.assertIn("Неделя 1 | созвоны", text)

    def test_pptx_slides_and_notes(self):
        from pptx import Presentation

        pr = Presentation()
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = "Пилот в пяти компаниях"
        s.notes_slide.notes_text_frame.text = "сказать про цену"
        p = self.dir / "deck.pptx"
        pr.save(p)
        text, info = docs.extract(p)
        self.assertEqual(info["slides"], 1)
        self.assertIn("[слайд 1]", text)
        self.assertIn("Заметки: сказать про цену", text)

    def test_csv_semicolon_cp1251(self):
        p = self.dir / "t.csv"
        p.write_bytes("Имя;Город\nОлег;Казань\n".encode("cp1251"))
        text, _ = docs.extract(p)
        self.assertIn("Олег | Казань", text)

    def test_odt_and_docx(self):
        odt = self.dir / "a.odt"
        with zipfile.ZipFile(odt, "w") as z:
            z.writestr(
                "content.xml",
                "<office:document-content><text:p>Первый абзац</text:p><text:p>Второй</text:p></office:document-content>",
            )
        text, _ = docs.extract(odt)
        self.assertIn("Первый абзац\nВторой", text)
        dx = self.dir / "a.docx"
        with zipfile.ZipFile(dx, "w") as z:
            z.writestr(
                "word/document.xml",
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body><w:p><w:r><w:t>Договор на пилот</w:t></w:r></w:p></w:body></w:document>",
            )
        self.assertIn("Договор на пилот", docs.extract(dx)[0])

    def test_html_drops_menus_and_scripts(self):
        p = self.dir / "page.html"
        p.write_text(
            "<html><head><title>Цены</title><script>var x=1</script></head><body><nav>Меню Главная</nav>"
            "<p>Тариф Про стоит 900 рублей</p><footer>© подвал</footer></body></html>",
            encoding="utf-8",
        )
        text, info = docs.extract(p)
        self.assertEqual(info["title"], "Цены")
        self.assertIn("Тариф Про стоит 900 рублей", text)
        self.assertNotIn("Меню", text)
        self.assertNotIn("var x", text)
        self.assertNotIn("подвал", text)

    def test_rtf_cyrillic(self):
        p = self.dir / "a.rtf"
        p.write_text(r"{\rtf1\ansi {\fonttbl\f0 Arial;}\f0 \'cf\'f0\'e8\'e2\'e5\'f2\par ok}", encoding="latin-1")
        self.assertIn("Привет", docs.extract(p)[0])

    def test_unknown_format_refused(self):
        p = self.dir / "a.xyz"
        p.write_bytes(b"\x00\x01")
        with self.assertRaises(ValueError):
            docs.extract(p)


class TestRetrieve(unittest.TestCase):
    def test_best_finds_the_relevant_part_of_a_long_text(self):
        filler = "\n\n".join(
            f"Раздел {i}. Общие положения о порядке работы и ответственности сторон." for i in range(300)
        )
        text = (
            "Договор оказания услуг.\n\n"
            + filler
            + "\n\nСтоимость пилота составляет 120 тысяч рублей в месяц.\n\n"
            + filler
        )
        out, cut = retrieve.best("сколько стоит пилот", text, budget=3000, head=200)
        self.assertTrue(cut)
        self.assertLessEqual(len(out), 3300)
        self.assertIn("120 тысяч", out)
        self.assertTrue(out.startswith("Договор оказания услуг"))

    def test_short_text_untouched(self):
        self.assertEqual(retrieve.best("x", "короткий текст"), ("короткий текст", False))

    def test_chunks_cut_text_without_spaces(self):
        parts = retrieve.chunks("а" * 5000)  # used to loop forever when no space to cut at
        self.assertGreater(len(parts), 4)


class FakeNet:
    """Replaces web._get: search pages and articles from a dict, no network."""

    def __init__(self, pages):
        self.pages, self.calls = pages, []

    def __call__(self, url, timeout=15, limit=3_000_000, direct=None, data=None, headers=None):
        self.calls.append(url)
        for k, v in self.pages.items():
            if k in url:
                return v, "text/html; charset=utf-8"
        raise OSError("нет сети в тесте")


DDG = (
    '<a class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fen.wikipedia.org%2Fwiki%2FLlama">llama.cpp сервер для локальных моделей</a>'
    '<a class="result__snippet" href="#">llama.cpp запускает локальные модели</a>'
    '<a class="result__a" href="https://docs.ggml.ai/gguf">Формат GGUF и llama.cpp</a>'
    '<a class="result__snippet" href="#">локальные модели в формате gguf</a>'
    '<a class="result__a" href="https://casino.example/">Онлайн казино</a><a class="result__snippet" href="#">бонус</a>'
)
ARTICLE = (
    "<html><title>llama.cpp</title><body><article>"
    + "<p>Вступление про историю проекта и авторов.</p>" * 60
    + "<p>llama.cpp запускает локальные модели на процессоре без видеокарты.</p>"
    + "<p>Прочие подробности сборки под разные платформы.</p>" * 60
    + "</article></body></html>"
)


def no_ddgs(query, backend):
    raise OSError("нет сети в тесте")


class TestWeb(unittest.TestCase):
    def setUp(self):
        self.orig, self.orig_d, web._cache = web._get, web._ddgs_text, {}
        web._get = self.net = FakeNet(
            {"duckduckgo.com/html": DDG, "en.wikipedia.org/wiki/Llama": ARTICLE, "docs.ggml.ai/gguf": ARTICLE}
        )
        web._ddgs_text = no_ddgs
        web.set_brave_key("")

    def tearDown(self):
        web._get, web._ddgs_text = self.orig, self.orig_d

    def test_junk_answer_is_skipped_for_the_next_engine(self):
        web._ddgs_text = lambda q, b: (
            [{"title": "Рецептор, активируемый пероксисомами", "url": "https://x.org/r", "snippet": ""}] * 3
        )
        res = web.search("рецепт борща")
        self.assertEqual(res, [])  # neither the junk nor anything from the fake DDG page

    def test_keywords_drop_question_words(self):
        self.assertEqual(web.keywords("Что такое llama.cpp и как его запустить?"), "llama.cpp запустить")
        self.assertEqual(
            web.keywords("Что такое формат GGUF у llama.cpp? Две-три фразы со ссылками."), "формат GGUF llama.cpp"
        )
        self.assertEqual(
            web.keywords("Найди, сколько весит faster-whisper medium. Коротко."), "весит faster-whisper medium"
        )

    def test_search_filters_junk_and_decodes_links(self):
        res = web.search("что такое llama.cpp локальные модели")
        urls = [r["url"] for r in res]
        self.assertEqual(urls, ["https://en.wikipedia.org/wiki/Llama", "https://docs.ggml.ai/gguf"])
        self.assertEqual(res[0]["engine"], "DuckDuckGo")

    def test_gather_sends_excerpts_not_pages(self):
        out = web.gather("llama.cpp локальные модели на процессоре", pages=2, budget=3000)
        self.assertEqual(len(out["sources"]), 2)
        self.assertIn("[1]", out["context"])
        self.assertIn("без видеокарты", out["context"])
        self.assertGreater(out["read_chars"], 3 * len(out["context"]))

    def test_nothing_found_is_empty_not_error(self):
        web._get = FakeNet({})
        orig_sleep = web.__dict__.get("_t")
        import time

        real = time.sleep
        time.sleep = lambda s: None
        try:
            self.assertEqual(web.search("абракадабра"), [])
        finally:
            time.sleep = real
        self.assertIsNone(orig_sleep)


class TestEconomy(Base):
    def setUp(self):
        super().setUp()
        self.v = Vault(self.dir / "v")
        init_vault(self.v)
        fill(self.v)

    def test_compact_memory_keeps_head_and_relevant_lines(self):
        head = "Кто: Иван, продакт.\n" * 60
        body = (
            "\n".join(f"Факт {i}: что-то про проект номер {i}." for i in range(400))
            + "\nОлег отвечает за SMS-рассылку."
        )
        out = EC.compact_memory(head + body, "кто делает SMS рассылку")
        self.assertTrue(out.startswith("Кто: Иван"))
        self.assertIn("Олег отвечает за SMS-рассылку", out)
        self.assertLess(len(out), 4600)

    def _chat_with_doc(self):
        big = (
            "Общие условия договора и порядок взаимодействия сторон.\n\n" * 400
            + "Штраф за просрочку оплаты — 0,1 процента в день.\n\n"
            + "Прочие условия.\n\n" * 200
        )
        f = attach.save(
            self.v, "c1", [{"name": "договор.txt", "data": __import__("base64").b64encode(big.encode()).decode()}]
        )
        return {
            "id": "c1",
            "messages": [
                {"role": "user", "text": "Вот договор", "files": f},
                {"role": "assistant", "text": "Прочитал."},
                {"role": "user", "text": "Какой штраф за просрочку оплаты?"},
            ],
        }, big

    def test_document_goes_by_relevant_passages(self):
        chat, big = self._chat_with_doc()
        msgs, full, sent = EC.build_history(chat, chat["messages"][-1]["text"])
        self.assertIn("0,1 процента в день", msgs[-1]["text"])
        self.assertIn("[Прикреплено: договор.txt", msgs[0]["text"])
        self.assertLess(sent, len(big) / 3)
        self.assertGreater(full, len(big))
        self.assertGreater(EC.savings(full, sent), 5000)

    def test_without_economy_the_whole_document_goes(self):
        chat, big = self._chat_with_doc()
        msgs, full, sent = EC.build_history(chat, chat["messages"][-1]["text"], economy=False)
        self.assertIn(big[:200], msgs[0]["text"])

    def test_old_turns_shortened_last_kept(self):
        long = "слово " * 300
        chat = {"messages": [{"role": "user" if i % 2 == 0 else "assistant", "text": f"{i} {long}"} for i in range(12)]}
        msgs, _, _ = EC.build_history(chat, "вопрос")
        self.assertIn("[сокращено]", msgs[0]["text"])
        self.assertNotIn("[сокращено]", msgs[-1]["text"])

    def test_local_prompt_is_the_same_every_turn(self):
        a = CH.system_prompt(self.v, "local", log=False, query="кто такой Олег")
        b = CH.system_prompt(self.v, "local", log=False, query="план пилота на неделю")
        self.assertEqual(a, b)  # the engine reuses what it already read

    def test_attach_reads_xlsx_into_text(self):
        from openpyxl import Workbook
        import base64
        import io

        wb = Workbook()
        wb.active.append(["Клиент", "Сумма"])
        wb.active.append(["Ромашка", 90000])
        buf = io.BytesIO()
        wb.save(buf)
        f = attach.save(self.v, "c2", [{"name": "клиенты.xlsx", "data": base64.b64encode(buf.getvalue()).decode()}])[0]
        self.assertEqual(f["kind"], "text")
        self.assertIn("Ромашка | 90000", Path(f["text_path"]).read_text(encoding="utf-8"))


class TestChatWeb(TestEconomy):
    def setUp(self):
        super().setUp()
        self.orig_s, self.orig_g, self.orig_d = PR.stream, web._get, web._ddgs_text
        web._ddgs_text = no_ddgs
        PR.stream = self.fake = FakeStream("По источнику [1] — работает на процессоре.")
        web._cache = {}
        web._get = FakeNet(
            {"duckduckgo.com/html": DDG, "en.wikipedia.org/wiki/Llama": ARTICLE, "docs.ggml.ai/gguf": ARTICLE}
        )

    def tearDown(self):
        PR.stream, web._get, web._ddgs_text = self.orig_s, self.orig_g, self.orig_d

    def test_tool_markup_never_reaches_the_owner(self):
        PR.stream = FakeStream(
            '<invoke name="Bash">\n<parameter name="command">ls</parameter>\n</invoke>\n\nФайлов не вижу.'
        )
        r = CH.send(self.v, "привет", model="claude_code:", sync=True, learn=False)
        self.assertEqual(CH.load_chat(self.v, r["chat"])["messages"][-1]["text"], "Файлов не вижу.")
        super().tearDown()

    def test_web_switch_adds_sources_and_saving(self):
        r = CH.send(
            self.v,
            "найди, что такое llama.cpp локальные модели",
            model="claude_code:",
            sync=True,
            web=True,
            learn=False,
        )
        c = CH.load_chat(self.v, r["chat"])
        ans = c["messages"][-1]
        self.assertEqual(
            [s["url"] for s in ans["sources"]], ["https://en.wikipedia.org/wiki/Llama", "https://docs.ggml.ai/gguf"]
        )
        self.assertGreater(ans["saved"], 0)
        last = self.fake.calls[0]["messages"][-1]["text"]
        self.assertIn("Выдержки из авторитетных источников", last)
        self.assertTrue(
            last.rstrip().endswith("найди, что такое llama.cpp локальные модели")
        )  # the question comes last
        self.assertTrue(c["messages"][0]["web"])

    def test_no_web_by_default(self):
        CH.send(self.v, "привет", model="claude_code:", sync=True, learn=False)
        self.assertNotIn("Выдержки из авторитетных источников", self.fake.calls[0]["messages"][-1]["text"])
        self.assertFalse(any("duckduckgo" in u for u in web._get.calls))

    def test_link_in_message_is_read_without_switch(self):
        CH.send(
            self.v,
            "перескажи https://en.wikipedia.org/wiki/Llama про процессор",
            model="claude_code:",
            sync=True,
            learn=False,
        )
        self.assertIn("без видеокарты", self.fake.calls[0]["messages"][-1]["text"])

    def test_economy_setting_off_sends_full_memory(self):
        P.save(self.v, {"economy": False})
        CH.send(self.v, "привет", model="claude_code:", sync=True, learn=False)
        self.assertIn("Иван", self.fake.calls[0]["system"])


class TestPhase(unittest.TestCase):
    def test_local_model_shows_reading_progress_and_time_left(self):
        g = {
            "text": "",
            "phase": "пишу",
            "pid": "local",
            "usage": {"progress": {"part": 0.64, "ms": 86170, "done": 2048, "total": 3200}},
        }
        self.assertEqual(CH._phase(g), "модель на компьютере читает вопрос и память · 64% · ещё около 45 с")
        g["sources"] = [{"url": "https://x"}]
        self.assertIn("читает вопрос и найденное", CH._phase(g))
        g["usage"] = {}
        self.assertEqual(CH._phase(g), "модель на компьютере читает вопрос и найденное")

    def test_thinking_is_hidden(self):
        self.assertEqual(CH.strip_thinking("<think>считаю 2+2</think>" + chr(10) + chr(10) + "Четыре."), "Четыре.")
        self.assertEqual(CH.strip_thinking("считаю…</think>Четыре."), "Четыре.")
        self.assertEqual(CH.strip_thinking("<think>ещё думаю"), "")
        g = {"text": "", "phase": "пишу", "pid": "local", "usage": {"reasoning": 120}}
        self.assertEqual(CH._phase(g), "модель рассуждает, потом ответит")

    def test_cloud_model_thinks_and_words_hide_the_phase(self):
        self.assertEqual(
            CH._phase({"text": "", "phase": "пишу", "pid": "claude_code", "who": "Claude"}), ""
        )  # plain chat: dots only
        self.assertEqual(
            CH._phase({"text": "", "phase": "пишу", "pid": "claude_code", "who": "Claude", "consortium": True}),
            "Claude думает",
        )
        self.assertEqual(CH._phase({"text": "Привет", "phase": "пишу", "pid": "local"}), "пишу")


class TestWhenToSearch(unittest.TestCase):
    def test_small_talk_and_self_dont_search(self):
        for t in (
            "Ну как дела?",
            "Привет",
            "Спасибо, отлично",
            "Что ты обо мне знаешь?",
            "Напиши письмо клиенту о задержке",
            "Что ты умеешь делать с pdf?",
            "Как в Mnema включить обучение?",
            "Что такое GGUF в llama.cpp?",
            "Кто такой Илон Маск?",
        ):  # general knowledge: the model knows it, the web only checks fresh data
            self.assertFalse(web.should_search(t), t)

    def test_outside_facts_search(self):
        for t in (
            "Какой курс доллара сегодня?",
            "Какая последняя версия Python?",
            "Сколько стоит iPhone 17?",
            "Найди отзывы о Notion",
        ):
            self.assertTrue(web.should_search(t), t)

    def test_only_serious_sources(self):
        for u in (
            "https://ru.wikipedia.org/wiki/X",
            "https://www.cbr.ru/",
            "https://docs.python.org/3/",
            "https://www.nalog.gov.ru/",
        ):
            self.assertTrue(web.trusted(u), u)
        for u in (
            "https://www.youtube.com/@x",
            "https://vkvideo.ru/x",
            "https://some-blog.ru/x",
            "https://www.kinopoisk.ru/series/1/",
        ):
            self.assertFalse(web.trusted(u), u)


if __name__ == "__main__":
    unittest.main()
